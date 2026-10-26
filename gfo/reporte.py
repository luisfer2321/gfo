"""Reporte de análisis financiero de la empresa en Excel.

Genera un libro con una hoja por análisis (estados financieros, horizontal,
vertical, razones, DuPont, punto de equilibrio, apalancamiento, flujo de
efectivo, presupuesto maestro y proforma). No incluye las aplicaciones
complementarias (presupuesto personal e inventario).

Los valores se calculan con las mismas funciones de ``gfo/calculos.py`` que usa
la aplicación, así que el Excel coincide con lo que se ve en pantalla.
"""

from __future__ import annotations

import datetime as dt
import io

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from gfo import calculos as calc
from gfo import datos

# Formatos numéricos con sus unidades
CS = '"C$" #,##0.00;[Red]-"C$" #,##0.00'
PCT = '0.00%;[Red]-0.00%'
VECES = '0.00 "veces";[Red]-0.00 "veces"'
DIAS = '#,##0 "días"'
NUM = '#,##0.00;[Red]-#,##0.00'
ENT = '#,##0;[Red]-#,##0'
FORMATO_UNIDAD = {"C$": CS, "%": PCT, "veces": VECES, "días": DIAS}

AZUL_OSCURO, AZUL, GRIS = "0D366B", "1C5CAB", "52514E"
SERIES = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "6250D6"]
EVALUACION = {"bueno": "Favorable", "aceptable": "Aceptable", "critico": "Desfavorable", "info": "Informativo"}
COLOR_EVALUACION = {"bueno": "006300", "aceptable": "9A6700", "critico": "B42828", "info": GRIS}

_delgada = Side(style="thin", color="C3C2B7")
_doble = Side(style="double", color="52514E")


class Hoja:
    """Escribe tablas con formato contable de arriba hacia abajo en una hoja."""

    def __init__(self, wb: Workbook, nombre: str, empresa: str, titulo: str, subtitulo: str = "", anchos=(46,)):
        self.ws = wb.create_sheet(nombre)
        self.ws.sheet_view.showGridLines = False
        self.ws["A1"] = empresa
        self.ws["A1"].font = Font(bold=True, size=14, color=AZUL_OSCURO)
        self.ws["A2"] = titulo
        self.ws["A2"].font = Font(bold=True, size=12, color=AZUL)
        if subtitulo:
            self.ws["A3"] = subtitulo
            self.ws["A3"].font = Font(italic=True, size=9, color=GRIS)
        self.fila = 5
        self.ws.column_dimensions["A"].width = anchos[0]
        for i in range(2, 16):
            self.ws.column_dimensions[get_column_letter(i)].width = anchos[i - 1] if i - 1 < len(anchos) else 17

    def seccion(self, texto: str) -> None:
        c = self.ws.cell(self.fila, 1, texto)
        c.font = Font(bold=True, size=11.5, color=AZUL_OSCURO)
        self.fila += 1

    def nota(self, texto: str, color: str = GRIS, negrita: bool = False) -> None:
        c = self.ws.cell(self.fila, 1, texto)
        c.font = Font(italic=not negrita, bold=negrita, size=9.5, color=color)
        self.fila += 1

    def espacio(self, n: int = 1) -> None:
        self.fila += n

    def tabla(self, encabezados: list[str], filas: list, fmt: str | list = CS) -> tuple[int, int]:
        """filas: (etiqueta, valores, estilo[, formato de la fila]); estilo en '', 'grupo', 'sub', 'total'.

        Un valor puede ser un número, un texto, None (se escribe N/D) o (número, formato).
        Devuelve (primera fila de datos, última fila).
        """
        ws = self.ws
        for j, h in enumerate(encabezados, start=1):
            c = ws.cell(self.fila, j, h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=AZUL_OSCURO)
            c.alignment = Alignment(horizontal="left" if j == 1 else "center", vertical="center", wrap_text=True)
        ws.row_dimensions[self.fila].height = 30 if any(len(str(h)) > 16 for h in encabezados) else 20
        self.fila += 1
        inicio = self.fila
        n = len(encabezados)
        for fila in filas:
            etiqueta, valores, estilo = fila[0], fila[1], fila[2]
            fmt_fila = fila[3] if len(fila) > 3 else None
            c = ws.cell(self.fila, 1, etiqueta)
            if estilo == "grupo":
                for j in range(1, n + 1):
                    ws.cell(self.fila, j).fill = PatternFill("solid", fgColor="EEF4FB")
                c.font = Font(bold=True, color=AZUL, size=9.5)
                self.fila += 1
                continue
            for j, v in enumerate(valores, start=2):
                formato = fmt_fila or (fmt[j - 2] if isinstance(fmt, list) else fmt)
                if isinstance(v, tuple):
                    v, formato = v
                if v is None or (isinstance(v, float) and (np.isnan(v) or np.isinf(v))):
                    celda = ws.cell(self.fila, j, "N/D")
                    celda.alignment = Alignment(horizontal="right")
                    celda.font = Font(color=GRIS)
                else:
                    celda = ws.cell(self.fila, j, float(v) if isinstance(v, (int, float, np.number)) else v)
                    if isinstance(v, (int, float, np.number)):
                        celda.number_format = formato
                    else:
                        celda.alignment = Alignment(wrap_text=True, vertical="top")
            for j in range(1, n + 1):
                celda = ws.cell(self.fila, j)
                celda.border = Border(bottom=Side(style="hair", color="E1E0D9"))
                if estilo in ("sub", "total"):
                    celda.font = Font(bold=True, color=celda.font.color)
                if estilo == "sub":
                    celda.fill = PatternFill("solid", fgColor="FAFAF8")
                if estilo == "total":
                    celda.border = Border(top=Side(style="thin", color=GRIS), bottom=_doble)
            self.fila += 1
        fin = self.fila - 1
        self.fila += 1
        return inicio, fin


def _pulir(g) -> None:
    """Título y leyenda fuera del área de trazado, ejes visibles y líneas rectas."""
    if g.title is not None:
        g.title.overlay = False
    g.legend.position = "b"
    g.legend.overlay = False
    g.x_axis.delete = False
    g.y_axis.delete = False
    for s in g.series:
        if getattr(s, "smooth", None) is not None or isinstance(g, (LineChart, ScatterChart)):
            s.smooth = False


def _grafico_barras(ws, titulo, cats, series_refs, ancla, alto=7.5, ancho=16, tipo="col", y_fmt=None):
    g = BarChart()
    g.type = tipo
    g.title = titulo
    g.style = 10
    g.height, g.width = alto, ancho
    g.grouping = "clustered"
    g.gapWidth = 60
    for i, ref in enumerate(series_refs):
        g.add_data(ref, titles_from_data=True)
    g.set_categories(cats)
    for i, s in enumerate(g.series):
        s.graphicalProperties.solidFill = SERIES[i % len(SERIES)]
        s.graphicalProperties.line.solidFill = SERIES[i % len(SERIES)]
    g.y_axis.majorGridlines.spPr = None
    if y_fmt:
        g.y_axis.numFmt = y_fmt
    _pulir(g)
    ws.add_chart(g, ancla)


def _evaluacion(clave, valor):
    return EVALUACION[calc.evaluar_razon(clave, valor)]


# ---------------------------------------------------------------------------
# Líneas de los estados
# ---------------------------------------------------------------------------

def _lineas_balance():
    lineas = []
    grupos = [("ACTIVO", ["Activo corriente", "Activo no corriente"], "activo_total"),
              ("PASIVO", ["Pasivo corriente", "Pasivo no corriente"], "pasivo_total"),
              ("PATRIMONIO", ["Patrimonio"], "patrimonio")]
    subtotal = {"Activo corriente": "activo_corriente", "Activo no corriente": "activo_no_corriente",
                "Pasivo corriente": "pasivo_corriente", "Pasivo no corriente": "pasivo_no_corriente"}
    for titulo, subgrupos, total in grupos:
        lineas.append((titulo, None, "grupo"))
        for g in subgrupos:
            lineas += [(n, k, "") for k, n, gg in datos.CUENTAS_BALANCE if gg == g]
            if g in subtotal:
                lineas.append((datos.NOMBRES[subtotal[g]], subtotal[g], "sub"))
        lineas.append((datos.NOMBRES[total], total, "total"))
    lineas.append(("TOTAL PASIVO + PATRIMONIO", "pasivo_patrimonio", "total"))
    return lineas


LINEAS_ER = [("ventas", ""), ("costo_ventas", "-"), ("utilidad_bruta", "sub"), ("gastos_venta", "-"),
             ("gastos_admin", "-"), ("depreciacion", "-"), ("gastos_operacion", "sub"), ("uaii", "sub"),
             ("intereses", "-"), ("otros_ingresos", ""), ("uai", "sub"), ("impuestos", "-"), ("utilidad_neta", "total")]


def _lineas_er():
    return [(("(−) " if e == "-" else "") + datos.NOMBRES[k], k, "" if e == "-" else e) for k, e in LINEAS_ER]


# ---------------------------------------------------------------------------
# Hojas
# ---------------------------------------------------------------------------

def generar(estados: pd.DataFrame, empresa: dict, P: dict, trimestres: pd.DataFrame,
            movimientos_caja: pd.DataFrame) -> bytes:
    """Construye el reporte completo y lo devuelve como bytes de un archivo .xlsx."""
    periodos = list(estados.columns)
    cuentas = {p: datos.cuentas_de(estados, p) for p in periodos}
    comp = {p: calc.periodo_completo(cuentas[p]) for p in periodos}
    razones = {p: calc.razones_financieras(cuentas[p]) for p in periodos}
    con_er = {p for p in periodos if calc.tiene_estado_resultados(cuentas[p])}
    claves_er = {k for k, _ in LINEAS_ER}

    def valor(p, k):
        """Valor de una cuenta; N/D para cifras de resultados en períodos que solo tienen balance."""
        return comp[p][k] if p in con_er or k not in claves_er else None

    ultimo = periodos[-1]
    anterior = periodos[-2] if len(periodos) > 1 else None
    nombre = empresa.get("nombre", "Empresa")
    hoy = dt.date.today().strftime("%d/%m/%Y")
    sub = f"Cifras en córdobas (C$) · Generado el {hoy} por el Sistema de Análisis y Planeación Financiera"

    pm = calc.presupuesto_maestro(datos.supuestos_de(P, cuentas[ultimo]), trimestres, cuentas[ultimo])
    proforma = calc.periodo_completo(pm["balance_proforma"])
    rz_pro = calc.razones_financieras(pm["balance_proforma"])
    base_cvu = datos.cvu_base_de(P, cuentas[ultimo])
    cvu = calc.cvu(base_cvu["precio"], base_cvu["cv_unitario"], base_cvu["costos_fijos"], base_cvu["unidades"])
    clas = calc.clasificar_costos(cuentas[ultimo])
    apal = calc.apalancamiento(comp[ultimo]["ventas"], clas["costos_variables"], clas["costos_fijos"],
                               comp[ultimo]["intereses"], P.get("ap_dividendos_pref", 0.0), cuentas[ultimo]["tasa_impuesto"])
    flujos = {(a, b): calc.flujo_efectivo_indirecto(cuentas[a], cuentas[b]) for a, b in zip(periodos, periodos[1:])}

    wb = Workbook()
    wb.remove(wb.active)

    # ------------------------------------------------------------ Resumen
    h = Hoja(wb, "Resumen", nombre, "Reporte de análisis y planeación financiera", sub, anchos=(40, 18, 18, 16, 60))
    h.nota(f"Actividad: {empresa.get('giro', '')}", color="0B0B0B")
    h.nota(f"Períodos analizados: {', '.join(periodos)} · Período principal: {ultimo}"
           + (f" · Comparado con: {anterior}" if anterior else ""), color="0B0B0B")
    h.nota("El presupuesto maestro y los estados proforma usan los supuestos registrados en «Planeación y presupuestos».")
    h.espacio()
    h.seccion("Contenido del reporte")
    for nombre_hoja, desc in [
        ("Balance General", "Balance General de cada período y ecuación contable"),
        ("Estado de Resultados", "Estado de Resultados y datos complementarios"),
        ("Análisis horizontal", "Variaciones absolutas y relativas entre períodos"),
        ("Análisis vertical", "Estructura porcentual del balance y del estado de resultados"),
        ("Razones financieras", "Liquidez, actividad, endeudamiento, rentabilidad y mercado"),
        ("Sistema DuPont", "Descomposición del ROE"),
        ("Punto de equilibrio", "Modelo costo-volumen-utilidad, escenarios y sensibilidad"),
        ("Apalancamiento", "GAO, GAF y GAT"),
        ("Flujo de efectivo", "Método indirecto y registro de entradas y salidas"),
        ("Presupuesto maestro", "Cédulas trimestrales y presupuesto de caja"),
        ("Estados proforma", "Estados proyectados e indicadores real vs. proforma"),
    ]:
        c = h.ws.cell(h.fila, 1, nombre_hoja)
        c.hyperlink = f"#'{nombre_hoja}'!A1"
        c.font = Font(color=AZUL, underline="single")
        h.ws.cell(h.fila, 2, desc).font = Font(color=GRIS)
        h.fila += 1
    h.espacio()
    h.espacio()
    h.seccion("Indicadores clave")

    def variacion(a, b):
        return calc.dividir(b - a, abs(a)) if None not in (a, b) and a else None

    indicadores = [
        ("Ventas netas", "ventas", CS, comp), ("Utilidad neta", "utilidad_neta", CS, comp),
        ("Total activo", "activo_total", CS, comp),
        ("Razón corriente", "razon_corriente", VECES, razones), ("Prueba ácida", "prueba_acida", VECES, razones),
        ("Razón de endeudamiento", "razon_deuda", PCT, razones), ("Margen neto", "margen_neto", PCT, razones),
        ("ROA", "roa", PCT, razones), ("ROE", "roe", PCT, razones),
    ]
    filas = []
    for etiqueta, k, f, fuente in indicadores:
        a = (valor(anterior, k) if fuente is comp else fuente[anterior][k]) if anterior else None
        b = valor(ultimo, k) if fuente is comp else fuente[ultimo][k]
        filas.append((etiqueta, [(a, f) if anterior else "", (b, f), (variacion(a, b), PCT) if anterior else "",
                                 _evaluacion(k, b) if fuente is razones else ""], ""))
    filas += [
        ("Punto de equilibrio (C$)", ["", (cvu["pe_monetario"], CS), "", f"{cvu['pe_unidades'] or 0:,.0f} unidades"], ""),
        ("Margen de seguridad", ["", (cvu["margen_seguridad_pct"], PCT), "", ""], ""),
        ("Apalancamiento total (GAT)", ["", (apal["gat"], VECES), "", f"GAO {apal['gao'] or 0:.2f} · GAF {apal['gaf'] or 0:.2f}"], ""),
    ]
    if anterior:
        f = flujos[(anterior, ultimo)]
        filas.append(("Flujo de efectivo de operación", ["", (f["flujo_operacion"], CS), "", ""], ""))
    filas += [
        ("Ventas presupuestadas (proforma)", ["", (proforma["ventas"], CS), (variacion(comp[ultimo]["ventas"], proforma["ventas"]), PCT), ""], ""),
        ("Utilidad neta proforma", ["", (proforma["utilidad_neta"], CS),
                                    (variacion(comp[ultimo]["utilidad_neta"], proforma["utilidad_neta"]), PCT), ""], ""),
        ("Saldo de caja proyectado", ["", (pm["saldos"][-1], CS), "", f"al cierre de {pm['trimestres'][-1]}"], ""),
    ]
    h.tabla(["Indicador", anterior or "", ultimo, "Variación", "Evaluación / detalle"], filas)

    h.seccion(f"Diagnóstico de las razones financieras · {ultimo}")
    diag = []
    for k, info in calc.INFO_RAZON.items():
        e = calc.evaluar_razon(k, razones[ultimo][k])
        if e != "info":
            diag.append((info["nombre"], [(razones[ultimo][k], FORMATO_UNIDAD.get(info["unidad"], NUM)), EVALUACION[e],
                                          calc.interpretar_razon(k, razones[ultimo][k])], ""))
    ini, fin = h.tabla(["Razón", "Valor", "Evaluación", "Interpretación"], diag)
    for r in range(ini, fin + 1):
        c = h.ws.cell(r, 3)
        clave = {v: k for k, v in EVALUACION.items()}.get(c.value, "info")
        c.font = Font(bold=True, color=COLOR_EVALUACION[clave])
        h.ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)

    h.seccion("Hallazgos para la toma de decisiones")
    for texto in _hallazgos(comp, razones, ultimo, anterior, proforma, rz_pro, pm, cvu, apal, flujos, P):
        h.nota("• " + texto, color="0B0B0B")
    h.espacio()

    h.seccion("Resultados por período (datos del gráfico)")
    cols = [p for p in periodos if p in con_er] + ["Proforma"]
    todos = {**comp, "Proforma": proforma}
    conceptos = [("Ventas", "ventas"), ("Utilidad bruta", "utilidad_bruta"), ("UAII", "uaii"), ("Utilidad neta", "utilidad_neta")]
    ini, fin = h.tabla(["Concepto"] + cols, [(c, [todos[p][k] for p in cols], "") for c, k in conceptos])
    _grafico_barras(h.ws, "Ventas y utilidades por período",
                    Reference(h.ws, min_col=1, min_row=ini, max_row=fin),
                    [Reference(h.ws, min_col=j, min_row=ini - 1, max_row=fin) for j in range(2, 2 + len(cols))],
                    f"A{h.fila + 1}", y_fmt='#,##0')
    h.fila += 17
    # ------------------------------------------------------------ Balance General
    h = Hoja(wb, "Balance General", nombre, "Balance General al cierre de cada período", sub)
    h.tabla(["Cuenta"] + periodos, [(n, [comp[p][k] for p in periodos] if k else [], e) for n, k, e in _lineas_balance()])
    h.tabla(["Comprobación"] + periodos, [("Activo − (Pasivo + Patrimonio)", [comp[p]["diferencia"] for p in periodos], "")])
    h.nota("La ecuación contable se cumple cuando la diferencia es cero.")

    # ------------------------------------------------------------ Estado de Resultados
    h = Hoja(wb, "Estado de Resultados", nombre, "Estado de Resultados de cada período", sub)
    h.tabla(["Cuenta"] + periodos, [(n, [valor(p, k) for p in periodos], e) for n, k, e in _lineas_er()])
    if len(con_er) < len(periodos):
        h.nota("N/D: " + ", ".join(p for p in periodos if p not in con_er) + " solo tiene Balance General.")
        h.espacio()
    h.seccion("Datos complementarios")
    h.tabla(["Dato"] + periodos, [
        ("Tasa de impuesto sobre la renta", [cuentas[p]["tasa_impuesto"] for p in periodos], "", PCT),
        ("Acciones en circulación", [cuentas[p]["acciones"] for p in periodos], "", ENT),
        ("Precio de mercado por acción", [cuentas[p]["precio_accion"] for p in periodos], "", CS),
        ("Dividendos pagados", [cuentas[p]["dividendos"] for p in periodos], "", CS),
    ])
    h.nota("Utilidad bruta = Ventas − Costo de ventas · UAII = Utilidad bruta − Gastos de operación · "
           "UAI = UAII − Intereses ± Otros ingresos (gastos) · Utilidad neta = UAI − Impuesto")

    # ------------------------------------------------------------ Análisis horizontal
    h = Hoja(wb, "Análisis horizontal", nombre, "Análisis horizontal", sub)
    if not anterior:
        h.nota("Se necesitan al menos dos períodos para el análisis horizontal.")
    for base, actual in zip(periodos, periodos[1:]):
        h.seccion(f"{base} → {actual}")
        for titulo, lineas in (("Balance General", _lineas_balance()), ("Estado de Resultados", _lineas_er())):
            filas = []
            for n, k, e in lineas:
                if not k:
                    filas.append((n, [], e))
                    continue
                b, a = valor(base, k), valor(actual, k)
                if None in (a, b):
                    filas.append((n, [b, a, None, (None, PCT)], e))
                else:
                    filas.append((n, [b, a, a - b, (calc.dividir(a - b, abs(b)) if b else None, PCT)], e))
            h.tabla([titulo, base, actual, "Variación C$", "Variación %"], filas)
    h.nota("Variación absoluta = Período actual − Período base · Variación relativa = Variación absoluta ÷ |Período base|")

    # ------------------------------------------------------------ Análisis vertical
    h = Hoja(wb, "Análisis vertical", nombre, "Análisis vertical", sub)
    h.seccion("Balance General (% del activo total)")
    h.tabla(["Cuenta"] + periodos, [(n, [calc.dividir(comp[p][k], comp[p]["activo_total"]) for p in periodos] if k else [], e)
                                    for n, k, e in _lineas_balance()], fmt=PCT)
    h.seccion("Estado de Resultados (% de las ventas netas)")
    h.tabla(["Cuenta"] + periodos, [(n, [calc.dividir(comp[p][k], comp[p]["ventas"]) for p in periodos], e)
                                    for n, k, e in _lineas_er()], fmt=PCT)
    h.nota("Balance: Cuenta ÷ Activo total · Estado de Resultados: Cuenta ÷ Ventas netas. N/D = el período no tiene ventas.")

    # ------------------------------------------------------------ Razones financieras
    h = Hoja(wb, "Razones financieras", nombre, "Razones financieras", sub, anchos=(34, 40) + (15,) * len(periodos) + (15, 70))
    filas, categoria = [], None
    for k, info in calc.INFO_RAZON.items():
        if info["categoria"] != categoria:
            categoria = info["categoria"]
            filas.append((categoria.upper(), [], "grupo"))
        f = FORMATO_UNIDAD.get(info["unidad"], NUM)
        filas.append((info["nombre"], [info["formula"]] + [(razones[p][k], f) for p in periodos]
                      + [_evaluacion(k, razones[ultimo][k]), calc.interpretar_razon(k, razones[ultimo][k])], ""))
    h.tabla(["Razón", "Fórmula"] + periodos + [f"Evaluación {ultimo}", f"Interpretación {ultimo}"], filas)
    h.seccion("Criterios de evaluación (referencias académicas generales)")
    h.tabla(["Razón", "Favorable", "Aceptable", "Desfavorable"], [
        ("Razón corriente", ["≥ 1.5", "1.0 – 1.5", "< 1.0"], ""), ("Prueba ácida", ["≥ 1.0", "0.7 – 1.0", "< 0.7"], ""),
        ("Razón de efectivo", ["≥ 0.2", "0.1 – 0.2", "< 0.1"], ""),
        ("Razón de endeudamiento", ["≤ 50 %", "50 % – 70 %", "> 70 %"], ""),
        ("Deuda a patrimonio", ["≤ 1.0", "1.0 – 2.0", "> 2.0"], ""), ("Cobertura de intereses", ["≥ 3.0", "1.5 – 3.0", "< 1.5"], ""),
        ("Margen bruto", ["≥ 25 %", "10 % – 25 %", "< 10 %"], ""), ("Margen operativo", ["≥ 10 %", "3 % – 10 %", "< 3 %"], ""),
        ("Margen neto", ["≥ 5 %", "0 % – 5 %", "< 0 %"], ""), ("ROA", ["≥ 5 %", "2 % – 5 %", "< 2 %"], ""),
        ("ROE", ["≥ 12 %", "5 % – 12 %", "< 5 %"], ""),
    ])
    h.nota("Las razones usan saldos al cierre. N/D = no se puede calcular (división entre cero o dato faltante).")
    h.ws.freeze_panes = "C6"

    # ------------------------------------------------------------ DuPont
    h = Hoja(wb, "Sistema DuPont", nombre, "Sistema DuPont", sub)
    dp = {p: calc.dupont(cuentas[p]) for p in periodos}
    h.tabla(["Factor"] + periodos, [
        ("Margen neto (Utilidad neta ÷ Ventas)", [(dp[p]["margen_neto"], PCT) for p in periodos], ""),
        ("(×) Rotación de activos (Ventas ÷ Activo total)", [(dp[p]["rotacion_activos"], VECES) for p in periodos], ""),
        ("(=) ROA", [(dp[p]["roa"], PCT) for p in periodos], "sub"),
        ("(×) Multiplicador del capital (Activo ÷ Patrimonio)", [(dp[p]["multiplicador_capital"], VECES) for p in periodos], ""),
        ("(=) ROE", [(dp[p]["roe"], PCT) for p in periodos], "total"),
        ("Comprobación: Utilidad neta ÷ Patrimonio", [(razones[p]["roe"], PCT) for p in periodos], ""),
    ])
    if anterior and dp[anterior]["roe"] is not None and dp[ultimo]["roe"] is not None:
        cambios = {n: calc.dividir((dp[ultimo][k] or 0) - (dp[anterior][k] or 0), dp[anterior][k]) or 0
                   for k, n in [("margen_neto", "el margen neto"), ("rotacion_activos", "la rotación de activos"),
                                ("multiplicador_capital", "el apalancamiento (multiplicador)")]}
        principal = max(cambios, key=lambda n: abs(cambios[n]))
        h.nota(f"El ROE pasó de {dp[anterior]['roe']:.2%} a {dp[ultimo]['roe']:.2%}; el factor que más cambió fue "
               f"{principal} ({cambios[principal]:+.1%}).", color="0B0B0B")
    h.nota("ROE = Margen neto × Rotación de activos totales × Multiplicador del capital")

    # ------------------------------------------------------------ Punto de equilibrio
    h = Hoja(wb, "Punto de equilibrio", nombre, f"Punto de equilibrio y modelo C-V-U · {ultimo}", sub,
             anchos=(46, 18, 34, 18))
    origen = P.get("pe_origen", "Estados financieros")
    h.seccion(f"Datos de entrada (origen: {origen.lower()})")
    h.tabla(["Dato", "Valor"], [
        ("Precio de venta unitario", [cvu["precio"]], "", CS),
        ("Costo variable unitario", [cvu["cv_unitario"]], "", CS),
        ("Costos fijos totales", [cvu["costos_fijos"]], "", CS),
        ("Unidades a vender", [cvu["unidades"]], "", NUM),
    ])
    if origen == "Estados financieros":
        h.nota(f"Del Estado de Resultados {ultimo}: unidades = Ventas ÷ Precio; costo variable unitario = (Costo de ventas + "
               "Gastos de venta) ÷ Unidades; costos fijos = Gastos de administración + Depreciación.")
        h.espacio()
    h.seccion("Resultados")
    h.tabla(["Concepto", "Valor", "Fórmula"], [
        ("Margen de contribución unitario", [(cvu["mcu"], CS), "Precio − Costo variable unitario"], ""),
        ("Razón de margen de contribución", [(cvu["rmc"], PCT), "MCu ÷ Precio"], ""),
        ("Punto de equilibrio (unidades)", [(cvu["pe_unidades"], NUM), "Costos fijos ÷ MCu"], "sub"),
        ("Punto de equilibrio (C$)", [(cvu["pe_monetario"], CS), "Costos fijos ÷ RMC"], "sub"),
        ("Ventas totales", [(cvu["ventas"], CS), "Unidades × Precio"], ""),
        ("(−) Costos variables", [(cvu["costos_variables"], CS), "Unidades × CVu"], ""),
        ("(=) Margen de contribución total", [(cvu["mc_total"], CS), "Ventas − Costos variables"], ""),
        ("(−) Costos fijos", [(cvu["costos_fijos"], CS), ""], ""),
        ("(=) Utilidad operativa", [(cvu["utilidad"], CS), "Unidades × MCu − Costos fijos"], "total"),
        ("Margen de seguridad (unidades)", [(cvu["margen_seguridad_u"], NUM), "Unidades − PE"], ""),
        ("Margen de seguridad (%)", [(cvu["margen_seguridad_pct"], PCT), "(Unidades − PE) ÷ Unidades"], ""),
        ("Grado de apalancamiento operativo", [(cvu["gao"], VECES), "MC total ÷ Utilidad operativa"], ""),
    ])
    if cvu["mcu"] > 0:
        esc = calc.escenario_cvu(cvu, P.get("esc_precio", 0) / 100, P.get("esc_cv", 0) / 100,
                                 P.get("esc_cf", 0) / 100, P.get("esc_vol", 0) / 100)
        h.seccion(f"Escenario simulado (precio {P.get('esc_precio', 0):+d} %, costo variable {P.get('esc_cv', 0):+d} %, "
                  f"costos fijos {P.get('esc_cf', 0):+d} %, volumen {P.get('esc_vol', 0):+d} %)")
        filas = []
        for n, k, f in [("Precio de venta unitario", "precio", CS), ("Costo variable unitario", "cv_unitario", CS),
                        ("Costos fijos", "costos_fijos", CS), ("Unidades vendidas", "unidades", NUM),
                        ("Punto de equilibrio (unidades)", "pe_unidades", NUM), ("Punto de equilibrio (C$)", "pe_monetario", CS),
                        ("Utilidad operativa", "utilidad", CS), ("Margen de seguridad", "margen_seguridad_pct", PCT)]:
            filas.append((n, [(cvu[k], f), (esc[k], f), (calc.dividir((esc[k] or 0) - (cvu[k] or 0), cvu[k]), PCT)], ""))
        h.tabla(["Concepto", "Escenario base", "Escenario simulado", "Cambio"], filas)
        h.seccion("Sensibilidad: efecto de un cambio de ±10 % en cada variable sobre la utilidad operativa")
        sens = []
        for n, args in [("Precio", (0.1, 0, 0, 0)), ("Costo variable", (0, 0.1, 0, 0)),
                        ("Costos fijos", (0, 0, 0.1, 0)), ("Volumen", (0, 0, 0, 0.1))]:
            sube = calc.escenario_cvu(cvu, *args)["utilidad"] - cvu["utilidad"]
            baja = calc.escenario_cvu(cvu, *[-a for a in args])["utilidad"] - cvu["utilidad"]
            sens.append((n, [sube, baja], ""))
        h.tabla(["Variable", "Variable +10 %", "Variable −10 %"], sens)
        obj = P.get("pe_utilidad_obj", 0.0)
        u = calc.unidades_para_utilidad(cvu["precio"], cvu["cv_unitario"], cvu["costos_fijos"], obj)
        h.seccion("Utilidad deseada")
        h.tabla(["Concepto", "Valor"], [
            ("Utilidad operativa deseada", [obj], "", CS),
            ("Unidades necesarias = (CF + Utilidad deseada) ÷ MCu", [u], "", NUM),
            ("Ventas necesarias", [u * cvu["precio"] if u else None], "", CS),
        ])
        # Datos y gráfico de la relación ingresos – costos
        h.seccion("Datos del gráfico de equilibrio")
        maximo = max(cvu["unidades"], cvu["pe_unidades"] or 0) * 1.35
        q = np.linspace(0, maximo, 11)
        ini, fin = h.tabla(["Unidades", "Ingresos totales", "Costos totales", "Costos fijos"],
                           [(round(x), [x * cvu["precio"], cvu["costos_fijos"] + x * cvu["cv_unitario"], cvu["costos_fijos"]], "")
                            for x in q])
        for r in range(ini, fin + 1):
            h.ws.cell(r, 1).number_format = ENT
        g = ScatterChart()
        g.title = "Ingresos vs. costos totales (eje X: unidades vendidas)"
        g.style = 13
        g.height, g.width = 9, 17
        xs = Reference(h.ws, min_col=1, min_row=ini, max_row=fin)
        for j, color in zip(range(2, 5), (SERIES[0], SERIES[1], "898781")):
            s = Series(Reference(h.ws, min_col=j, min_row=ini - 1, max_row=fin), xs, title_from_data=True)
            s.marker.symbol = "none"
            s.graphicalProperties.line.solidFill = color
            s.graphicalProperties.line.width = 28000
            s.smooth = False
            g.series.append(s)
        g.x_axis.majorGridlines = None
        g.x_axis.numFmt = "#,##0"
        g.y_axis.numFmt = "#,##0"
        _pulir(g)
        h.ws.add_chart(g, f"F{ini - 1}")
    else:
        h.nota("El precio no supera al costo variable unitario: no existe punto de equilibrio.", color="B42828")

    # ------------------------------------------------------------ Apalancamiento
    h = Hoja(wb, "Apalancamiento", nombre, "Grados de apalancamiento operativo, financiero y total", sub)
    con_ventas = [p for p in periodos if comp[p]["ventas"]]
    ap = {}
    for p in con_ventas:
        c = calc.clasificar_costos(cuentas[p])
        ap[p] = calc.apalancamiento(comp[p]["ventas"], c["costos_variables"], c["costos_fijos"], comp[p]["intereses"],
                                    P.get("ap_dividendos_pref", 0.0), cuentas[p]["tasa_impuesto"])
    if ap:
        h.tabla(["Concepto"] + con_ventas, [
            ("Ventas", [ap[p]["ventas"] for p in con_ventas], ""),
            ("(−) Costos variables (costo de ventas + gastos de venta)", [ap[p]["costos_variables"] for p in con_ventas], ""),
            ("(=) Margen de contribución", [ap[p]["mc"] for p in con_ventas], "sub"),
            ("(−) Costos fijos (administración + depreciación)", [ap[p]["costos_fijos"] for p in con_ventas], ""),
            ("(=) Utilidad operativa (UAII)", [ap[p]["uaii"] for p in con_ventas], "sub"),
            ("(−) Gastos por intereses", [ap[p]["intereses"] for p in con_ventas], ""),
            ("(=) Utilidad antes de impuestos", [ap[p]["uai"] for p in con_ventas], "total"),
            ("GAO = Margen de contribución ÷ UAII", [(ap[p]["gao"], VECES) for p in con_ventas], ""),
            ("GAF = UAII ÷ (UAII − Intereses − DP ÷ (1 − t))", [(ap[p]["gaf"], VECES) for p in con_ventas], ""),
            ("GAT = GAO × GAF", [(ap[p]["gat"], VECES) for p in con_ventas], "sub"),
        ])
        cambio = P.get("ap_cambio_ventas", 10) / 100
        u = con_ventas[-1]
        ef = calc.efecto_apalancamiento(ap[u]["gao"] or 0, ap[u]["gaf"] or 0, cambio)
        h.seccion(f"Efecto de una variación de {cambio:+.0%} en las ventas · {u}")
        h.tabla(["Concepto", "Variación"], [
            ("Ventas", [ef["cambio_ventas"]], "", PCT),
            ("Utilidad operativa (UAII) = GAO × Δ% ventas", [ef["cambio_uaii"]], "", PCT),
            ("Utilidad por acción (UPA) = GAT × Δ% ventas", [ef["cambio_upa"]], "", PCT),
        ])
    else:
        h.nota("Ningún período tiene Estado de Resultados con ventas.")

    # ------------------------------------------------------------ Flujo de efectivo
    h = Hoja(wb, "Flujo de efectivo", nombre, "Estado de flujo de efectivo", sub, anchos=(58, 18, 18, 18))
    if not flujos:
        h.nota("El método indirecto necesita dos balances consecutivos.")
    for (a, b), f in flujos.items():
        h.seccion(f"Método indirecto · {a} → {b}")
        filas = [("ACTIVIDADES DE OPERACIÓN", [], "grupo")] + [(n, [v], "") for n, v in f["operacion"]]
        filas += [("Flujo neto de operación", [f["flujo_operacion"]], "sub"), ("ACTIVIDADES DE INVERSIÓN", [], "grupo")]
        filas += [(n, [v], "") for n, v in f["inversion"]] + [("Flujo neto de inversión", [f["flujo_inversion"]], "sub")]
        filas += [("ACTIVIDADES DE FINANCIAMIENTO", [], "grupo")] + [(n, [v], "") for n, v in f["financiamiento"]]
        filas += [("Flujo neto de financiamiento", [f["flujo_financiamiento"]], "sub"),
                  ("Variación neta del efectivo", [f["variacion_neta"]], "sub"),
                  (f"(+) Efectivo al inicio ({a})", [f["saldo_inicial"]], ""),
                  (f"Efectivo al final ({b})", [f["saldo_final"]], "total")]
        h.tabla(["Concepto", "C$"], filas)
        h.nota(("Conciliado: coincide con el efectivo del Balance General " if f["conciliado"] else
                "No concilia con el efectivo del Balance General ") + f"{b} (C$ {f['saldo_balance']:,.2f}).",
               color="006300" if f["conciliado"] else "B42828")
        h.espacio()
    h.seccion("Registro de entradas y salidas de efectivo")
    mov = movimientos_caja.dropna(subset=["Monto"])
    if len(mov):
        h.tabla(["Concepto", "Actividad", "Tipo", "Monto"],
                [(r["Concepto"], [r["Actividad"], r["Tipo"], (r["Monto"], CS)], "") for _, r in mov.iterrows()])
        d = calc.flujo_efectivo_directo(mov, P.get("fd_saldo_inicial", 0.0))
        filas = [(r["Actividad"], [r["Entradas"], r["Salidas"], r["Flujo neto"]], "") for _, r in d["resumen"].iterrows()]
        filas += [("Variación neta", ["", "", d["variacion_neta"]], "sub"),
                  ("Saldo inicial", ["", "", d["saldo_inicial"]], ""), ("Saldo final", ["", "", d["saldo_final"]], "total")]
        h.tabla(["Actividad", "Entradas", "Salidas", "Flujo neto"], filas)

    # ------------------------------------------------------------ Presupuesto maestro
    h = Hoja(wb, "Presupuesto maestro", nombre, f"Presupuesto maestro trimestral (saldos iniciales del Balance {ultimo})", sub,
             anchos=(64,))
    s = datos.supuestos_de(P, cuentas[ultimo])
    h.seccion("Supuestos de planificación")
    h.tabla(["Supuesto", "Valor"], [
        ("Precio de venta unitario", [s["precio"]], "", CS), ("Costo unitario de compra", [s["costo_unitario"]], "", CS),
        ("Ventas al contado", [s["pct_contado"]], "", PCT), ("Compras al contado", [s["pct_compras_contado"]], "", PCT),
        ("Inventario final deseado (% ventas del trimestre siguiente)", [s["pct_inventario_final"]], "", PCT),
        ("Gastos de venta (% de las ventas)", [s["pct_gastos_venta"]], "", PCT),
        ("Gastos de administración por trimestre", [s["gastos_admin_trim"]], "", CS),
        ("Depreciación por trimestre", [s["depreciacion_trim"]], "", CS),
        ("Intereses anuales", [s["intereses_anuales"]], "", CS), ("Saldo mínimo de caja", [s["saldo_minimo"]], "", CS),
    ])
    unidades = {"Unidades a vender", "(+) Inventario final deseado", "(=) Necesidad total", "(−) Inventario inicial",
                "(=) Unidades a comprar"}
    for titulo, clave in [("Presupuesto de ventas", "ventas"), ("Cédula de cobros", "cobros"),
                          ("Presupuesto de compras", "compras"), ("Cédula de pagos a proveedores", "pagos"),
                          ("Costo de bienes vendidos", "cbv"), ("Gastos de operación", "gastos"),
                          ("Presupuesto de caja", "caja")]:
        df = pm[clave]
        h.seccion(titulo)
        filas = []
        for idx, fila in df.iterrows():
            estilo = "total" if idx.startswith(("Total", "(=) Costo", "Saldo final")) else "sub" if idx.startswith("(=)") else ""
            filas.append((idx, [None if pd.isna(v) else v for v in fila], estilo, NUM if idx in unidades else CS))
        ini, fin = h.tabla(["Concepto"] + list(df.columns), filas)
        for r in range(ini, fin + 1):
            for j in range(2, len(df.columns) + 2):
                if h.ws.cell(r, j).value == "N/D":
                    h.ws.cell(r, j).value = ""
        if clave == "caja":
            fila_caja = {h.ws.cell(r, 1).value: r for r in range(ini, fin + 1)}
    n_t = len(pm["trimestres"])
    cats = Reference(h.ws, min_col=2, max_col=1 + n_t, min_row=ini - 1)
    g = BarChart()
    g.title, g.style, g.height, g.width = "Cobros vs. egresos por trimestre", 10, 7.5, 15
    for nombre_fila, color in (("(+) Cobros a clientes", SERIES[0]), ("Total egresos", SERIES[1])):
        r = fila_caja[nombre_fila]
        g.add_data(Reference(h.ws, min_col=1, max_col=1 + n_t, min_row=r), from_rows=True, titles_from_data=True)
        g.series[-1].graphicalProperties.solidFill = color
        g.series[-1].graphicalProperties.line.solidFill = color
    g.set_categories(cats)
    g.y_axis.numFmt = "#,##0"
    _pulir(g)
    h.ws.add_chart(g, f"H{ini}")
    g = LineChart()
    g.title, g.style, g.height, g.width = "Saldo de efectivo proyectado", 12, 7.5, 15
    for nombre_fila, color in (("Saldo antes de financiamiento", SERIES[1]), ("Saldo final", SERIES[0])):
        g.add_data(Reference(h.ws, min_col=1, max_col=1 + n_t, min_row=fila_caja[nombre_fila]), from_rows=True,
                   titles_from_data=True)
        g.series[-1].graphicalProperties.line.solidFill = color
        g.series[-1].graphicalProperties.line.width = 28000
    g.set_categories(cats)
    g.y_axis.numFmt = "#,##0"
    _pulir(g)
    h.ws.add_chart(g, f"H{ini + 18}")

    # ------------------------------------------------------------ Estados proforma
    h = Hoja(wb, "Estados proforma", nombre, "Estados financieros proforma (año presupuestado)", sub)
    real = comp[ultimo]
    h.seccion("Estado de Resultados proforma")
    h.tabla(["Cuenta", ultimo, "Proforma", "Cambio"],
            [(n, [real[k], proforma[k], (calc.dividir(proforma[k] - real[k], abs(real[k])) if real[k] else None, PCT)], e)
             for n, k, e in _lineas_er()])
    h.seccion("Balance General proforma")
    h.tabla(["Cuenta", ultimo, "Proforma"], [(n, [real[k], proforma[k]] if k else [], e) for n, k, e in _lineas_balance()])
    h.nota("Activo − (Pasivo + Patrimonio) proforma: " + f"C$ {proforma['diferencia']:,.2f}")
    h.espacio()
    h.seccion("Indicadores clave: real vs. proforma")
    claves = ["razon_corriente", "prueba_acida", "capital_trabajo", "dias_inventario", "periodo_cobro", "razon_deuda",
              "cobertura_intereses", "margen_bruto", "margen_neto", "roa", "roe"]
    h.tabla(["Indicador"] + periodos + ["Proforma", "Evaluación proforma"],
            [(calc.INFO_RAZON[k]["nombre"], [(razones[p][k], FORMATO_UNIDAD.get(calc.INFO_RAZON[k]["unidad"], NUM)) for p in periodos]
              + [(rz_pro[k], FORMATO_UNIDAD.get(calc.INFO_RAZON[k]["unidad"], NUM)), _evaluacion(k, rz_pro[k])], "")
             for k in claves])

    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = AZUL if ws.title == "Resumen" else "86B6EF"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToHeight = 0

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _hallazgos(comp, razones, ultimo, anterior, proforma, rz_pro, pm, cvu, apal, flujos, P) -> list[str]:
    """Conclusiones breves para la toma de decisiones."""
    t = []
    d = comp[ultimo]
    if anterior and comp[anterior]["ventas"] and d["ventas"]:
        cv = (d["ventas"] - comp[anterior]["ventas"]) / comp[anterior]["ventas"]
        cu = calc.dividir(d["utilidad_neta"] - comp[anterior]["utilidad_neta"], abs(comp[anterior]["utilidad_neta"]))
        t.append(f"Las ventas variaron {cv:+.2%} entre {anterior} y {ultimo}"
                 + (f" y la utilidad neta {cu:+.2%}." if cu is not None else "."))
    r = razones[ultimo]
    for k in ("razon_corriente", "razon_deuda", "roe"):
        e = calc.evaluar_razon(k, r[k])
        if e != "info":
            t.append(f"{calc.INFO_RAZON[k]['nombre']}: {EVALUACION[e].lower()}. {calc.interpretar_razon(k, r[k])}")
    if cvu["pe_unidades"]:
        t.append(f"El punto de equilibrio es de {cvu['pe_unidades']:,.0f} unidades (C$ {cvu['pe_monetario']:,.2f}); "
                 f"las ventas pueden caer hasta {cvu['margen_seguridad_pct']:.2%} antes de generar pérdidas.")
    if apal["gat"] is not None:
        t.append(f"Con un apalancamiento total de {apal['gat']:.2f}, un cambio de 10 % en las ventas cambia la utilidad "
                 f"por acción en {apal['gat'] * 10:.1f} %.")
    if anterior:
        f = flujos[(anterior, ultimo)]
        t.append(f"La operación generó C$ {f['flujo_operacion']:,.2f} de efectivo; "
                 + ("cubre las inversiones del período." if f["flujo_operacion"] >= -f["flujo_inversion"]
                    else "las inversiones requirieron financiamiento adicional."))
    fin = pm["caja"].loc["Financiamiento (+) / pago (−)", pm["trimestres"]]
    if fin.max() > 0:
        t.append(f"El presupuesto de caja requiere C$ {fin.max():,.2f} de financiamiento de corto plazo en {fin.idxmax()}"
                 + (f"; quedan C$ {pm['deuda_nueva']:,.2f} pendientes al cierre." if pm["deuda_nueva"] > 0 else "; se paga dentro del año."))
    else:
        t.append("El efectivo proyectado se mantiene sobre el saldo mínimo durante todo el año.")
    if d["utilidad_neta"]:
        t.append(f"La utilidad neta proforma es C$ {proforma['utilidad_neta']:,.2f} "
                 f"({(proforma['utilidad_neta'] - d['utilidad_neta']) / abs(d['utilidad_neta']):+.2%} respecto de {ultimo}).")
    return t
