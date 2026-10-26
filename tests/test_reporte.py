"""Pruebas del reporte de análisis en Excel."""

import io

import openpyxl
import pytest

from gfo import datos, reporte


@pytest.fixture(scope="module")
def libro():
    contenido = reporte.generar(datos._estados_df(datos.ESTADOS_EJEMPLO), datos.EMPRESA, dict(datos.PARAMETROS_EJEMPLO),
                                datos.TRIMESTRES_EJEMPLO, datos.MOVIMIENTOS_CAJA_EJEMPLO)
    return openpyxl.load_workbook(io.BytesIO(contenido))


def buscar(ws, etiqueta):
    """Valores de la primera fila cuya columna A es la etiqueta."""
    for fila in ws.iter_rows(values_only=True):
        if fila[0] == etiqueta:
            return fila[1:]
    raise AssertionError(f"No se encontró la fila {etiqueta!r} en {ws.title}")


def test_hojas_del_reporte(libro):
    assert libro.sheetnames == ["Resumen", "Balance General", "Estado de Resultados", "Análisis horizontal",
                                "Análisis vertical", "Razones financieras", "Sistema DuPont", "Punto de equilibrio",
                                "Apalancamiento", "Flujo de efectivo", "Presupuesto maestro", "Estados proforma"]


def test_no_incluye_aplicaciones_complementarias(libro):
    texto = " ".join(str(c) for ws in libro.worksheets for fila in ws.iter_rows(values_only=True) for c in fila if c)
    assert "inventario de productos" not in texto.lower()
    assert "presupuesto personal" not in texto.lower()
    assert "Arroz" not in texto and "Salario mensual" not in texto


def test_valores_coinciden_con_los_calculos(libro):
    assert buscar(libro["Balance General"], "TOTAL ACTIVO")[:2] == (6_000_000, 6_600_000)
    assert buscar(libro["Estado de Resultados"], "Utilidad neta")[1] == pytest.approx(729_400)
    roe = buscar(libro["Razones financieras"], "Rendimiento sobre patrimonio (ROE)")
    assert roe[2] == pytest.approx(729_400 / 3_450_000)
    assert buscar(libro["Sistema DuPont"], "(=) ROE")[1] == pytest.approx(729_400 / 3_450_000)
    assert buscar(libro["Flujo de efectivo"], "Variación neta del efectivo")[0] == pytest.approx(70_000)
    assert buscar(libro["Análisis horizontal"], "Ventas netas")[3] == pytest.approx(0.15)


def test_tiene_graficos(libro):
    assert len(libro["Resumen"]._charts) == 1
    assert len(libro["Punto de equilibrio"]._charts) == 1
    assert len(libro["Presupuesto maestro"]._charts) == 2
